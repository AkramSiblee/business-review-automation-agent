#!/usr/bin/env python3
"""
Builds Business_Review_Complex_Input_FY2025.xlsx: a synthetic FY2025 group (Halcyon Industrial Holdings,
8 entities, 7 currencies) that plants every scenario in
.claude/skills/business-review-automation/references/complex-scenarios-playbook.md, and EXPECTED.json with the
reference values a correct review cycle should reproduce. README.md in this folder is the answer key.

  python sample-data/complex-test-data/generate_complex_dataset.py

Deterministic (fixed seed): re-running produces identical data. Every entity's ledger balances by construction
(interim trial-balance convention, as in the Meridian sample: retained earnings carried at the period-opening
balance, current-period profit still in the P&L accounts); the deliberate traps are applied only at export.
"""
import json
import os
import random
from collections import defaultdict
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(HERE, "Business_Review_Complex_Input_FY2025.xlsx")
OUT_JSON = os.path.join(HERE, "EXPECTED.json")
random.seed(20251231)

FY = 2025
M = [f"{FY}-{m:02d}" for m in range(1, 13)]
PREV = {M[0]: "2024-12", **{M[i]: M[i - 1] for i in range(1, 12)}}
MNUM = {p: int(p[5:]) for p in M}


def eom(p):
    y, m = int(p[:4]), int(p[5:])
    return date(y, m, [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1])


# ---------------------------------------------------------------- FX (units of USD per 1 unit of currency)
def drift(a, b, wobble):
    return [round(a + (b - a) * i / 12 + random.uniform(-wobble, wobble), 6) for i in range(13)]


KEYS = ["2024-12"] + M
SPOT = {
    "EUR": dict(zip(KEYS, drift(1.035, 1.100, 0.006))),
    "GBP": dict(zip(KEYS, drift(1.250, 1.300, 0.006))),
    "BRL": dict(zip(KEYS, drift(0.1620, 0.1830, 0.0015))),
    "MXN": dict(zip(KEYS, drift(0.0490, 0.0545, 0.0004))),
    "INR": dict(zip(KEYS, drift(0.01170, 0.01140, 0.00005))),
    "CAD": dict(zip(KEYS, drift(0.695, 0.730, 0.004))),
}
t = [0.02830, 0.02780, 0.02740]  # TRY: steady slide, a 20% step devaluation in March, then 2.5% a month
t.append(round(t[-1] * 0.80, 6))
while len(t) < 13:
    t.append(round(t[-1] * 0.975, 6))
SPOT["TRY"] = dict(zip(KEYS, t))
AVG = {c: {p: round((s[PREV[p]] + s[p]) / 2, 6) for p in M} for c, s in SPOT.items()}
for c in AVG:
    AVG[c]["2024-12"] = SPOT[c]["2024-12"]
AVG["USD"] = {p: 1.0 for p in KEYS}
SPOT["USD"] = {p: 1.0 for p in KEYS}

# ---------------------------------------------------------------- chart of accounts
A = {  # code: (type, english name)
    "1000": ("Asset", "Cash & Equivalents"), "1100": ("Asset", "Accounts Receivable"),
    "1150": ("Asset", "Allowance for Doubtful Accounts"), "1200": ("Asset", "Inventory"),
    "1300": ("Asset", "Prepaid Expenses"), "1400": ("Asset", "Property, Plant & Equipment, net"),
    "1450": ("Asset", "Right-of-Use Assets"), "1460": ("Asset", "Capitalized Development Costs"),
    "1500": ("Asset", "Investment in Subsidiaries"), "1510": ("Asset", "Investment in Joint Venture"),
    "1600": ("Asset", "Intercompany Receivable"), "1900": ("Asset", "Other Assets"),
    "2000": ("Liability", "Accounts Payable"), "2050": ("Liability", "Intercompany Payable"),
    "2100": ("Liability", "Accrued Liabilities"), "2150": ("Liability", "Warranty Provision"),
    "2200": ("Liability", "Deferred Revenue"), "2300": ("Liability", "Lease Liabilities"),
    "2400": ("Liability", "Revolving Credit Facility"), "2450": ("Liability", "Term Loan"),
    "2500": ("Liability", "Contingent Consideration (Earn-out)"), "2900": ("Liability", "Other Liabilities"),
    "3000": ("Equity", "Share Capital"), "3100": ("Equity", "Retained Earnings"),
    "4000": ("Revenue", "Product Revenue"), "4100": ("Revenue", "Service Revenue"),
    "4900": ("Revenue", "Intercompany Sales"), "4910": ("Revenue", "Intercompany Management Fee Income"),
    "5000": ("COGS", "Materials & Production"), "5100": ("COGS", "Direct Labor"),
    "5200": ("COGS", "Warranty Expense"), "5900": ("COGS", "Intercompany Purchases"),
    "6000": ("OpEx-S&M", "Sales & Marketing Salaries"), "6100": ("OpEx-S&M", "Marketing Programs"),
    "6110": ("OpEx-S&M", "Digital Marketing"), "6200": ("OpEx-R&D", "R&D Salaries"),
    "6300": ("OpEx-G&A", "G&A Salaries"), "6400": ("OpEx-G&A", "Rent & Facilities"),
    "6500": ("OpEx-G&A", "Depreciation & Amortization"), "6600": ("OpEx-G&A", "Professional Fees"),
    "6700": ("OpEx-G&A", "Intercompany Management Fee"), "7000": ("Other Income", "Interest Income"),
    "7100": ("Other Expense", "Interest Expense"), "7200": ("Other Income", "FX Gain/(Loss)"),
    "7300": ("Other Income", "Gain on Disposal of Subsidiary"), "7900": ("Other Expense", "Restructuring"),
    "7910": ("Other Expense", "Acquisition Deal Costs"), "7920": ("Other Expense", "Earn-out Fair Value Change"),
    "8000": ("Other Expense", "Income Tax"),
}
TYPE = {c: v[0] for c, v in A.items()}
INCOME = {"Revenue", "Other Income"}
DEPT = {"5100": "Operations", "6000": "Sales", "6100": "Marketing", "6110": "Marketing", "6200": "Engineering",
        "6300": "G&A", "6400": "G&A", "6500": "G&A", "6600": "G&A", "6700": "G&A"}
LOCAL = {
    "DE": {"1000": "Kasse & Bankguthaben", "1100": "Forderungen aus Lieferungen und Leistungen", "1200": "Vorräte",
           "1300": "Aktive Rechnungsabgrenzung", "1400": "Sachanlagen", "1450": "Nutzungsrechte (IFRS 16)",
           "2000": "Verbindlichkeiten aus Lieferungen und Leistungen",
           "2050": "Verbindlichkeiten gegenüber verbundenen Unternehmen", "2100": "Sonstige Rückstellungen",
           "2150": "Gewährleistungsrückstellung", "2200": "Passive Rechnungsabgrenzung",
           "2300": "Leasingverbindlichkeiten", "3000": "Gezeichnetes Kapital", "3100": "Gewinnvortrag",
           "4000": "Umsatzerlöse Produkte", "4100": "Umsatzerlöse Dienstleistungen", "5000": "Materialaufwand",
           "5100": "Löhne Lager & Logistik", "5200": "Gewährleistungsaufwand",
           "5900": "Wareneinkauf verbundene Unternehmen", "6000": "Gehälter Vertrieb & Marketing",
           "6100": "Werbeaufwand", "6300": "Gehälter Verwaltung", "6400": "Mietaufwand", "6500": "Abschreibungen",
           "6600": "Rechts- und Beratungskosten", "7100": "Zinsaufwand", "7900": "Restrukturierungsaufwand",
           "8000": "Ertragsteuern"},
    "PT": {"1000": "Caixa e Equivalentes", "1100": "Contas a Receber", "1200": "Estoques",
           "1300": "Despesas Antecipadas", "1400": "Imobilizado", "1450": "Direito de Uso", "2000": "Fornecedores",
           "2100": "Provisões e Encargos", "2200": "Receita Diferida", "2300": "Passivo de Arrendamento",
           "3000": "Capital Social", "3100": "Lucros Acumulados", "4000": "Receita de Produtos",
           "4100": "Receita de Serviços", "5000": "Custo dos Materiais", "5100": "Mão de Obra Direta",
           "6000": "Salários Vendas e Marketing", "6100": "Marketing", "6200": "Salários P&D",
           "6300": "Salários Administrativos", "6400": "Aluguel", "6500": "Depreciação e Amortização",
           "6600": "Honorários Profissionais", "6700": "Taxa de Administração Intercompany",
           "8000": "IRPJ/CSLL"},
    "ES": {"1000": "Efectivo y Equivalentes", "1100": "Cuentas por Cobrar", "1200": "Inventarios",
           "1300": "Pagos Anticipados", "1400": "Propiedad, Planta y Equipo", "1450": "Activos por Derecho de Uso",
           "1900": "Otros Activos", "2000": "Proveedores", "2100": "Pasivos Acumulados", "2200": "Ingresos Diferidos",
           "2300": "Pasivos por Arrendamiento", "3000": "Capital Social", "3100": "Resultados Acumulados",
           "4000": "Ingresos por Productos", "4100": "Ingresos por Servicios", "5000": "Costo de Materiales",
           "5100": "Mano de Obra Directa", "6000": "Sueldos Ventas y Marketing", "6100": "Mercadotecnia",
           "6200": "Sueldos Ingeniería", "6300": "Sueldos Administración", "6400": "Renta e Instalaciones",
           "6500": "Depreciación y Amortización", "6600": "Honorarios Profesionales", "7100": "Gastos Financieros",
           "8000": "ISR"},
}
LANG = {"DE01": "DE", "BR01": "PT", "MX01": "ES"}


def acct_name(e, code):
    if e == "HQ01" and code == "1200":
        return "Inventory (LIFO)"
    return LOCAL.get(LANG.get(e, ""), {}).get(code, A[code][1])


# ---------------------------------------------------------------- headcount rosters (drive the salary lines)
TITLES = {"Operations": ["Production Associate", "Warehouse Lead", "Quality Technician"],
          "Engineering": ["Engineer", "Senior Engineer", "Engineering Manager"],
          "Sales": ["Account Executive", "Sales Manager", "Inside Sales Rep"],
          "Marketing": ["Marketing Specialist", "Product Marketing Manager"],
          "G&A": ["Accountant", "HR Business Partner", "Office Manager", "Controller"]}
DEPT_ACCT = {"Operations": "5100", "Engineering": "6200", "Sales": "6000", "Marketing": "6000", "G&A": "6300"}
FMT = {"MDY": "%m/%d/%Y", "DMY": "%d/%m/%Y", "DOT": "%d.%m.%Y"}
ROSTER = []  # rows for the 03_Headcount tab


def roster(e, ccy, country, plan, fmt, extra=(), all_leave=None):
    n = 0
    for dept, count, cost, cur, ctry in [(d, c, k, ccy, country) for d, (c, k) in plan.items()] + list(extra):
        for _ in range(count):
            n += 1
            hire = date(random.randint(2014, 2024), random.randint(1, 12), random.randint(1, 28))
            if random.random() < 0.08:
                hire = date(FY, random.randint(1, 10), random.randint(1, 28))
            term, ttype = None, "N-A"
            if random.random() < 0.07:
                m = random.randint(max(2, hire.month + 1 if hire.year == FY else 2), 12)
                term, ttype = date(FY, m, random.randint(1, 28)), random.choice(["Voluntary", "Voluntary", "Involuntary"])
            if all_leave and (term is None or term > all_leave):
                term, ttype = all_leave, "Transfer (disposal)"
            ROSTER.append({"e": e, "id": f"{e}-{n:04d}", "dept": dept, "title": random.choice(TITLES[dept]),
                           "hire": hire, "term": term, "ttype": ttype, "fte": 0.5 if random.random() < 0.04 else 1.0,
                           "cost": round(cost * random.uniform(0.85, 1.15), 2), "ccy": cur, "country": ctry,
                           "fmt": fmt})


roster("HQ01", "USD", "United States", {"Operations": (48, 6500), "Engineering": (40, 12000), "Sales": (26, 10000),
                                        "Marketing": (8, 9000), "G&A": (24, 9500)}, "MDY",
       extra=[("Engineering", 5, 13000, "CAD", "Canada")])
roster("DE01", "EUR", "Germany", {"Operations": (22, 4600), "Sales": (20, 6500), "Marketing": (5, 6000),
                                  "G&A": (13, 6200)}, "DOT")
roster("BR01", "BRL", "Brazil", {"Operations": (40, 7000), "Sales": (18, 14000), "Marketing": (5, 12000),
                                 "G&A": (15, 13000), "Engineering": (4, 18000)}, "DMY")
roster("TR01", "TRY", "Turkey", {"Operations": (25, 75000), "Sales": (12, 110000), "G&A": (10, 105000),
                                 "Engineering": (3, 140000)}, "ISO")
roster("MX01", "MXN", "Mexico", {"Operations": (22, 28000), "Sales": (10, 45000), "G&A": (8, 42000),
                                 "Engineering": (5, 55000)}, "DMY")
roster("UK01", "GBP", "United Kingdom", {"Operations": (10, 3800), "Sales": (10, 5500), "Marketing": (3, 5000),
                                         "G&A": (6, 5200)}, "DMY", all_leave=date(FY, 9, 30))

SAL = defaultdict(lambda: defaultdict(float))  # (entity, account) -> period -> native cost of staff employed that month
for r in ROSTER:
    for p in M:
        if r["hire"] <= eom(p) and (r["term"] is None or r["term"] >= eom(p).replace(day=1)):
            v = r["cost"] * (AVG[r["ccy"]][p] if r["ccy"] == "CAD" else 1.0)  # CAD staff are paid via HQ01 (USD)
            SAL[(r["e"], DEPT_ACCT[r["dept"]])][p] += v


def sal(e, acct, p, factor=1.0):
    return SAL[(e, acct)][p] * factor


# ---------------------------------------------------------------- ledger engine
def simulate(months, cap, opening, step):
    """Balances by construction: cash = liabilities + capital + RE(opening) + NI - non-cash assets."""
    assets0 = sum(v for k, v in opening.items() if TYPE[k] == "Asset")
    re_open = round(assets0 - sum(v for k, v in opening.items() if TYPE[k] == "Liability") - cap, 2)
    prev, ni_prev, res = dict(opening), None, []
    for i, p in enumerate(months):
        if i:
            re_open = round(re_open + ni_prev, 2)
        pnl, bal, cap = step(i, p, prev, cap)
        pnl = {k: round(v, 2) for k, v in pnl.items() if abs(v) >= 0.005}
        bal = {k: round(v, 2) for k, v in bal.items()}
        ni = round(sum(v if TYPE[k] in INCOME else -v for k, v in pnl.items()), 2)
        non_cash = sum(v for k, v in bal.items() if TYPE[k] == "Asset" and k != "1000")
        liab = sum(v for k, v in bal.items() if TYPE[k] == "Liability")
        bal["1000"] = round(liab + cap + re_open + ni - non_cash, 2)
        res.append({"period": p, "pnl": pnl, "bs": bal, "cap": cap, "re_open": re_open, "ni": ni})
        prev, ni_prev = bal, ni
    return res


def pretax(pnl):
    return sum(v if TYPE[k] in INCOME else -v for k, v in pnl.items() if k != "8000")


def with_tax(pnl, rate):
    pnl["8000"] = max(0.0, rate * pretax(pnl))
    return pnl


# ---------------------------------------------------------------- entities
HQ_IC = {p: 550_000 * 1.004 ** i for i, p in enumerate(M)}  # HQ01 sale of goods to DE01, USD
CAPDEV = {p: 150_000 if MNUM[p] <= 6 else 400_000 for p in M}
INV_SUBS0 = {"DE01": 8_800_000, "BR01": 2_800_000, "TR01": 3_500_000, "UK01": 6_000_000}
MX_CONSIDERATION, MX_CASH, MX_EARNOUT = 25_000_000, 22_000_000, 3_000_000
UK_PROCEEDS, UK_COST = 9_000_000, INV_SUBS0["UK01"]
RESTRUCT_MONTHS = (2, 5, 8, 11)


def hq01(term_loan):
    def step(i, p, prev, cap):
        mth = MNUM[p]
        prod = 3_200_000 * 1.006 ** i * (1.35 if mth == 12 else 1.0)  # December: bill-and-hold uplift
        svc = 800_000 * 0.997 ** i
        revolver_prev = prev.get("2400", 0.0)
        pnl = {"4000": prod, "4100": svc, "4900": HQ_IC[p], "4910": 50_000,
               "5000": 0.55 * (prod + HQ_IC[p]), "5100": sal("HQ01", "5100", p),
               "6000": sal("HQ01", "6000", p), ("6100" if mth <= 6 else "6110"): 150_000 if mth <= 6 else 160_000,
               "6200": sal("HQ01", "6200", p) - CAPDEV[p],  # part of engineering payroll capitalized
               "6300": sal("HQ01", "6300", p), "6400": 120_000, "6500": 140_000, "6600": 90_000,
               "7000": 15_000, "7100": term_loan * 0.0525 / 12 + revolver_prev * 0.07 / 12,
               "7200": random.uniform(-25_000, 25_000)}
        if mth in RESTRUCT_MONTHS:
            pnl["7900"] = 120_000
        if mth == 7:
            pnl["7910"] = 650_000
        if mth == 9:
            pnl["7300"] = UK_PROCEEDS - UK_COST
        if mth == 12:
            pnl["7920"] = 300_000
        with_tax(pnl, 0.21)
        dso = 60 if mth == 12 else 45
        dio = {**{m: 70 for m in range(1, 7)}, 7: 85, 8: 95, 9: 100, 10: 108, 11: 115, 12: 120}[mth]
        dpo = {10: 60, 11: 70, 12: 80}.get(mth, 45)
        ar = (prod + svc) * dso / 30
        invest = sum(INV_SUBS0.values()) + (MX_CONSIDERATION if mth >= 7 else 0) - (UK_COST if mth >= 9 else 0)
        salaries = sum(sal("HQ01", a, p) for a in ("5100", "6000", "6200", "6300"))
        bal = {"1100": ar, "1150": -0.02 * ar, "1200": (pnl["5000"] + pnl["5100"]) * dio / 30, "1300": 250_000,
               "1400": prev["1400"] + 200_000 - 120_000, "1450": prev["1450"] - 40_000,
               "1460": prev["1460"] + CAPDEV[p] - 20_000, "1500": invest, "1510": 4_000_000,
               "1600": HQ_IC[p] + HQ_IC.get(PREV[p], HQ_IC[M[0]]), "1900": 500_000,
               "2000": (pnl["5000"] + pnl.get("6100", 0) + pnl.get("6110", 0) + 120_000 + 90_000) * dpo / 30,
               "2100": 0.5 * salaries, "2200": svc * (1.6 - 0.8 * i / 11), "2300": prev["2300"] - 38_000,
               "2400": 5_000_000 if mth >= 11 else 0.0, "2450": term_loan,
               "2500": (MX_EARNOUT if mth >= 7 else 0.0) + (300_000 if mth == 12 else 0.0), "2900": 300_000}
        return pnl, bal, cap

    opening = {"1000": 32_000_000, "1100": 6_500_000, "1150": -130_000, "1200": 4_300_000, "1300": 250_000,
               "1400": 18_000_000, "1450": 2_400_000, "1460": 1_200_000, "1500": sum(INV_SUBS0.values()),
               "1510": 4_000_000, "1600": 1_100_000, "1900": 500_000, "2000": 2_000_000, "2100": 800_000,
               "2200": 1_300_000, "2300": 2_500_000, "2400": 0.0, "2450": term_loan, "2500": 0.0, "2900": 300_000}
    return simulate(M, 25_000_000, opening, step)


def de01():
    def step(i, p, prev, cap):
        mth = MNUM[p]
        prod, svc = 1_900_000 * 1.003 ** i, 250_000
        icp = HQ_IC[p] / AVG["EUR"][p] - (150_000 if mth == 12 else 0)  # December: goods in transit not booked
        accrual = 0.01 * (prod + svc)
        release = 600_000 if mth == 11 else 0  # undisclosed warranty-provision release
        claims = 0.008 * (prod + svc)
        f = 1.21  # GL carries employer social charges the HR roster excludes
        pnl = {"4000": prod, "4100": svc, "5000": 0.48 * prod, "5100": sal("DE01", "5100", p, f),
               "5200": accrual - release, "5900": icp, "6000": sal("DE01", "6000", p, f), "6100": 60_000,
               "6300": sal("DE01", "6300", p, f), "6400": 70_000, "6500": 40_000, "6600": 30_000, "7100": 4_000}
        if mth == 3:
            pnl["7900"] = 400_000
        with_tax(pnl, 0.30)
        bal = {"1100": (prod + svc) * 50 / 30, "1200": (pnl["5000"] + icp) * 60 / 30, "1300": 50_000,
               "1400": prev["1400"] - 5_000, "1450": prev["1450"] - 25_000, "2000": pnl["5000"] * 40 / 30,
               "2050": icp + (prev["_icp"] if "_icp" in prev else icp),
               "2100": 0.4 * sum(sal("DE01", a, p, f) for a in ("5100", "6000", "6300")),
               "2150": prev["2150"] + accrual - claims - release, "2200": 0.5 * svc, "2300": prev["2300"] - 24_000}
        return pnl, bal, cap

    opening = {"1000": 4_000_000, "1100": 3_500_000, "1200": 2_200_000, "1300": 50_000, "1400": 3_000_000,
               "1450": 1_200_000, "2000": 800_000, "2050": 1_000_000, "2100": 300_000, "2150": 900_000,
               "2200": 120_000, "2300": 1_250_000}
    res = []

    def wrapped(i, p, prev, cap):  # carry last month's IC purchases for the 2-month payable
        prev = dict(prev)
        if res_icp:
            prev["_icp"] = res_icp[-1]
        pnl, bal, cap = step(i, p, prev, cap)
        res_icp.append(pnl["5900"])
        return pnl, bal, cap

    res_icp = []
    res = simulate(M, 8_000_000, opening, wrapped)
    return res


def br01():
    def step(i, p, prev, cap):
        prod, svc = 5_200_000 * 1.004 ** i, 800_000
        pnl = {"4000": prod, "4100": svc, "5000": 0.60 * prod, "5100": sal("BR01", "5100", p),
               "6000": sal("BR01", "6000", p), "6100": 120_000, "6200": sal("BR01", "6200", p),
               "6300": sal("BR01", "6300", p), "6400": 150_000, "6500": 110_000, "6600": 60_000,
               "6700": 50_000 / AVG["BRL"][p]}
        with_tax(pnl, 0.34)
        bal = {"1100": (prod + svc) * 60 / 30, "1200": pnl["5000"] * 70 / 30, "1300": 300_000,
               "1400": prev["1400"] - 10_000, "1450": prev["1450"] - 15_000, "2000": pnl["5000"] * 45 / 30,
               "2100": 0.5 * sum(sal("BR01", a, p) for a in ("5100", "6000", "6200", "6300")),
               "2200": 0.75 * svc, "2300": prev["2300"] - 14_000}
        return pnl, bal, cap

    opening = {"1000": 9_000_000, "1100": 9_000_000, "1200": 7_000_000, "1300": 300_000, "1400": 12_000_000,
               "1450": 2_000_000, "2000": 4_000_000, "2100": 1_500_000, "2200": 600_000, "2300": 2_100_000}
    return simulate(M, 20_000_000, opening, step)


TR_USD_PAYABLE = 4_000_000  # USD-denominated supplier balance carried by TR01


def tr01():
    def step(i, p, prev, cap):
        rev_i, cost_i = 1.035 ** i, 1.045 ** i  # nominal: costs inflate faster than prices
        prod, svc = 32_000_000 * rev_i, 6_000_000 * rev_i
        fx_loss = TR_USD_PAYABLE * (1 / SPOT["TRY"][p] - 1 / SPOT["TRY"][PREV[p]])
        pnl = {"4000": prod, "4100": svc, "5000": 0.80 * 32_000_000 * cost_i, "5100": sal("TR01", "5100", p, cost_i),
               "6000": sal("TR01", "6000", p, cost_i), "6200": sal("TR01", "6200", p, cost_i),
               "6300": sal("TR01", "6300", p, cost_i), "6400": 1_500_000 * cost_i, "6500": 1_200_000,
               "6600": 400_000 * cost_i, "7200": -fx_loss}
        with_tax(pnl, 0.25)
        bal = {"1100": (prod + svc) * 50 / 30, "1200": pnl["5000"] * 60 / 30,
               "1400": prev["1400"] + 500_000 - 1_200_000,
               "2000": pnl["5000"] * 40 / 30 + TR_USD_PAYABLE / SPOT["TRY"][p],
               "2100": 0.4 * sum(pnl[a] for a in ("5100", "6000", "6200", "6300"))}
        return pnl, bal, cap

    opening = {"1000": 90_000_000, "1100": 55_000_000, "1200": 40_000_000, "1400": 90_000_000,
               "2000": 26_000_000 + TR_USD_PAYABLE / SPOT["TRY"]["2024-12"], "2100": 10_000_000}
    return simulate(M, 150_000_000, opening, step)


MX_OPEN = {"1000": 40_000_000, "1100": 30_000_000, "1200": 25_000_000, "1300": 2_000_000, "1400": 180_000_000,
           "1450": 15_000_000, "1900": 10_000_000, "2000": 20_000_000, "2100": 8_000_000, "2200": 5_000_000,
           "2300": 16_000_000}
MX_CAP = 200_000_000


def mx01():
    months = M[6:]

    def step(i, p, prev, cap):
        prod, svc = 15_000_000 * 1.005 ** i, 3_000_000
        pnl = {"4000": prod, "4100": svc, "5000": 0.62 * prod, "5100": sal("MX01", "5100", p),
               "6000": sal("MX01", "6000", p), "6100": 400_000, "6200": sal("MX01", "6200", p),
               "6300": sal("MX01", "6300", p), "6400": 900_000, "6500": 700_000, "6600": 300_000, "7100": 120_000}
        with_tax(pnl, 0.30)
        bal = {"1100": (prod + svc) * 50 / 30, "1200": pnl["5000"] * 55 / 30, "1300": 2_000_000,
               "1400": prev["1400"] + 400_000 - 600_000, "1450": prev["1450"] - 100_000, "1900": 10_000_000,
               "2000": pnl["5000"] * 40 / 30,
               "2100": 0.5 * sum(sal("MX01", a, p) for a in ("5100", "6000", "6200", "6300")),
               "2200": 1.6 * svc, "2300": prev["2300"] - 90_000}
        return pnl, bal, cap

    return simulate(months, MX_CAP, dict(MX_OPEN), step)


def uk01():
    months = M[:9]

    def step(i, p, prev, cap):
        prod, svc = 950_000 * 1.002 ** i, 150_000
        pnl = {"4000": prod, "4100": svc, "5000": 0.58 * prod, "5100": sal("UK01", "5100", p),
               "6000": sal("UK01", "6000", p), "6100": 40_000, "6300": sal("UK01", "6300", p), "6400": 45_000,
               "6500": 25_000, "6600": 20_000}
        with_tax(pnl, 0.25)
        bal = {"1100": (prod + svc) * 45 / 30, "1200": pnl["5000"] * 45 / 30, "1400": prev["1400"] - 5_000,
               "2000": pnl["5000"] * 35 / 30, "2100": 0.4 * sum(sal("UK01", a, p) for a in ("5100", "6000", "6300")),
               "2200": 2.0 * svc}
        return pnl, bal, cap

    opening = {"1000": 2_500_000, "1100": 1_800_000, "1200": 600_000, "1400": 900_000, "2000": 500_000,
               "2100": 200_000, "2200": 300_000}
    return simulate(months, 3_000_000, opening, step)


def jv01():  # 40% joint venture: management-accounts P&L only, no balance sheet supplied
    out = []
    for i, p in enumerate(M):
        pnl = {"4000": 120_000_000 * 1.01 ** i, "5000": 60_000_000 * 1.01 ** i, "6300": 30_000_000,
               "6500": 5_000_000}
        with_tax(pnl, 0.25)
        pnl = {k: round(v, 2) for k, v in pnl.items()}
        out.append({"period": p, "pnl": pnl, "bs": {}, "cap": 0, "re_open": 0,
                    "ni": round(sum(v if TYPE[k] in INCOME else -v for k, v in pnl.items()), 2)})
    return out


def hld1():
    return simulate(M, 50_000, {"1000": 50_000}, lambda i, p, prev, cap: ({}, {}, cap))


# ---------------------------------------------------------------- consolidated covenant figures, and the term loan
CONSOL = {"HQ01": "USD", "DE01": "EUR", "BR01": "BRL", "TR01": "TRY", "MX01": "MXN", "UK01": "GBP"}
EBITDA_EXCL = {"4900", "4910", "5900", "6700", "6500"}  # intercompany and D&A


def group_ebitda(ledgers, periods):
    tot = 0.0
    for e, ccy in CONSOL.items():
        for r in ledgers[e]:
            if r["period"] in periods:
                v = sum((x if TYPE[k] == "Revenue" else -x) for k, x in r["pnl"].items()
                        if TYPE[k] in ("Revenue", "COGS") or TYPE[k].startswith("OpEx") and k not in EBITDA_EXCL
                        if k not in EBITDA_EXCL)
                tot += v * AVG[ccy][r["period"]]
    return tot


def group_cash(ledgers, p, exclude_restricted=True):
    tot = 0.0
    for e, ccy in list(CONSOL.items()) + [("HLD1", "USD")]:
        r = next((x for x in ledgers[e] if x["period"] == p), None)
        if r:
            tot += r["bs"]["1000"] * SPOT[ccy][p]
    return tot - (HQ_RESTRICTED if exclude_restricted else 0)


def interest(ledgers, periods):
    return sum(r["pnl"].get("7100", 0) * AVG[CONSOL[e]][r["period"]] for e in CONSOL for r in ledgers[e]
               if r["period"] in periods)


HQ_RESTRICTED, HQ_EUR_ACCOUNT = 1_500_000, 300_000
COVENANT_LEVERAGE = 1.50
TARGET_LEVERAGE = 1.40  # deliberately inside 15% of the covenant
state = random.getstate()
term_loan = 50_000_000.0
for _ in range(6):  # interest on the loan moves cash slightly, so iterate to a fixed point
    random.setstate(state)
    L = {"HQ01": hq01(term_loan), "DE01": de01(), "BR01": br01(), "TR01": tr01(), "MX01": mx01(),
         "UK01": uk01(), "JV01": jv01(), "HLD1": hld1()}
    ebitda = group_ebitda(L, M)
    revolver = L["HQ01"][-1]["bs"]["2400"]
    target = round((TARGET_LEVERAGE * ebitda - revolver + group_cash(L, M[-1])) / 100_000) * 100_000
    if abs(target - term_loan) < 1:
        break
    term_loan = target

for e, led in L.items():
    for r in led:
        assert r["bs"].get("1000", 0) >= 0, f"{e} {r['period']} cash negative"
        if r["bs"]:
            eq = sum(v for k, v in r["bs"].items() if TYPE[k] == "Asset") -                 sum(v for k, v in r["bs"].items() if TYPE[k] == "Liability")
            assert eq > 0, f"{e} {r['period']} equity not positive"
assert sum(r["ni"] for r in L["TR01"]) < 0, "TR01 must be loss-making (impairment trigger)"

# ---------------------------------------------------------------- export helpers
wb = openpyxl.Workbook()
wb.remove(wb.active)
HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F3864")


def sheet(name, headers, rows, pre=()):
    ws = wb.create_sheet(name)
    for line in pre:
        ws.append(list(line))
    ws.append(headers)
    hr = ws.max_row
    for c in ws[hr]:
        c.font, c.fill = HDR, FILL
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(40, len(str(h)) + 4))
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    return ws


NOTES = {("HQ01", "7900"): "Plant footprint optimisation (one-time)",
         ("HQ01", "7910"): "Advisory fees - acquisition of MX01",
         ("HQ01", "7300"): "Sale of UK01 - proceeds less cost of investment (parent books)",
         ("HQ01", "1510"): "Investment in JV01 carried at cost", ("HQ01", "2450"): "Term loan, matures 30 Jun 2026",
         ("TR01", "7200"): "Revaluation of USD-denominated supplier payables",
         ("HQ01", "2500"): "Earn-out for MX01 at fair value"}
ONE_TIME = {("HQ01", "7900"), ("HQ01", "7910"), ("HQ01", "7300"), ("DE01", "7900")}


def tb_lines(e, led, ytd=False):
    """Yield (period, code, name, type, dept, debit, credit, note). YTD exports carry cumulative P&L and
    retained earnings as at the start of the year (which is how such ERP exports still balance)."""
    cum = defaultdict(float)
    for r in led:
        p = r["period"]
        for code, v in sorted(r["bs"].items()):
            t = TYPE[code]
            dr, cr = ((v, 0) if v >= 0 else (0, -v)) if t == "Asset" else ((0, v) if v >= 0 else (-v, 0))
            yield p, code, t, dr, cr
        if r["bs"]:
            yield p, "3000", "Equity", 0, r["cap"]
            re = led[0]["re_open"] if ytd else r["re_open"]
            yield p, "3100", "Equity", *((0, re) if re >= 0 else (-re, 0))
        for code, v in r["pnl"].items():
            cum[code] += v
        # a YTD export lists every account posted so far this year, not just this month's
        for code in sorted(cum if ytd else r["pnl"]):
            v = round(cum[code], 2) if ytd else r["pnl"][code]
            t = TYPE[code]
            if t in INCOME:
                yield p, code, t, *((0, v) if v >= 0 else (-v, 0))
            else:
                yield p, code, t, *((v, 0) if v >= 0 else (0, -v))


CCY = {"HQ01": "USD", "DE01": "EUR", "BR01": "BRL", "TR01": "TRY", "MX01": "MXN", "UK01": "GBP", "JV01": "INR",
       "HLD1": "USD"}


def gl_row(e, p, code, t, dr, cr):
    note = NOTES.get((e, code), "")
    if (e, code) == ("HQ01", "2400") and p == "2025-11":
        note = "Drawn on revolving credit facility"
    if (e, code) == ("HQ01", "1500") and p == "2025-07":
        note = "Includes investment in MX01 (acquired 1 Jul 2025)"
    return [e, p, code, acct_name(e, code), t, DEPT.get(code, ""), round(dr, 2), round(cr, 2), CCY[e],
            "Y" if (e, code) in ONE_TIME and (dr or cr) else "N", note]


# ---------------------------------------------------------------- 00 config
cfg = [["HQ01", "Halcyon Industrial Holdings Inc.", "United States", "USD", "US GAAP", "January", 100, "Full", "USD",
        "Parent company; inventory carried on LIFO"],
       ["DE01", "Halcyon Vertrieb GmbH", "Germany", "EUR", "IFRS", "January", 100, "Full", "USD", ""],
       ["BR01", "Halcyon do Brasil Ltda.", "Brazil", "BRL", "IFRS", "January", "70%", "Full", "USD",
        "30% held by a local partner"],
       ["TR01", "Halcyon Endustri A.S.", "Turkey", "TRY", "IFRS", "January", 100, "Full", "USD", ""],
       ["MX01", "Halcyon Mexico S.A. de C.V.", "Mexico", "MXN", "IFRS", "January", 100, "Full", "USD",
        "Acquired 1 Jul 2025"],
       ["UK01", "Halcyon UK Ltd", "United Kingdom", "GBP", "IFRS", "January", 100, "Full", "USD", "Sold 30 Sep 2025"],
       ["JV01", "Halcyon Saraswati Precision Pvt Ltd", "India", "INR", "Ind AS (Local GAAP)", "April", 0.4, "Equity",
        "USD", "Joint venture; management accounts (P&L only) supplied"],
       ["HLD1", "Halcyon Holdings (Cayman) Ltd", "Cayman Islands", "USD", "US GAAP", "January", 100, "Full", "USD",
        "Dormant holding company"]]
sheet("00_Company_Config", ["Entity Code", "Entity Name", "Country", "Functional Currency", "Accounting Standard",
                            "Fiscal Year Start Month", "Ownership %", "Consolidation Method", "Reporting Currency",
                            "Notes"], cfg)

# ---------------------------------------------------------------- 01 GL (main tab)
GLH = ["Entity Code", "Period", "Account Code", "Account Name", "Account Type", "Department", "Debit", "Credit",
       "Currency", "One-Time Item", "Notes"]
rows = []
for e in ("HQ01", "DE01", "TR01", "UK01", "JV01", "HLD1"):
    lines = list(tb_lines(e, L[e], ytd=(e == "DE01")))
    for p, code, t, dr, cr in lines:
        rows.append(gl_row(e, p, code, t, dr, cr))
    if e == "TR01":  # subtotal rows pasted in from a formatted report
        for r in L["TR01"]:
            rev = r["pnl"]["4000"] + r["pnl"]["4100"]
            cogs = r["pnl"]["5000"] + r["pnl"]["5100"]
            rows.append(["TR01", r["period"], None, "Total Revenue", "Revenue", "", 0, round(rev, 2), "TRY", "N", ""])
            rows.append(["TR01", r["period"], None, "Gross Profit", "", "", 0, round(rev - cogs, 2), "TRY", "N", ""])
rows += [["HQ01", "2025-13", "6600", "Professional Fees", "OpEx-G&A", "G&A", 85_000, 0, "USD", "N",
          "Audit adjustment - year-end legal accrual"],
         ["HQ01", "2025-13", "2100", "Accrued Liabilities", "Liability", "", 0, 85_000, "USD", "N",
          "Audit adjustment - year-end legal accrual"],
         ["HQ01", "2026-02", "4000", "Product Revenue", "Revenue", "", 0, 3_450_000, "USD", "N", "FC"],
         ["HQ01", "2026-02", "5000", "Materials & Production", "COGS", "", 1_450_000, 0, "USD", "N", "FC"]]
sheet("01_GL_TrialBalance", GLH, rows)

# BR01: Portuguese names, reported in BRL thousands, April missing from the export
br_rows = [[e_, p, c, n, t, d, round(dr / 1000, 3), round(cr / 1000, 3), "BRL", o, nt]
           for e_, p, c, n, t, d, dr, cr, _, o, nt in
           (gl_row("BR01", p, code, t, dr, cr) for p, code, t, dr, cr in tb_lines("BR01", L["BR01"]))
           if p != "2025-04"]
sheet("BR01 GL (BRL 000)", GLH, br_rows,
      pre=[["Halcyon do Brasil Ltda. - Balancete de Verificacao FY2025"], ["Valores em BRL '000"], []])

# MX01: Spanish names, single signed Amount column (debit +, credit -), post-acquisition months only
mx_rows = [["MX01", p, code, acct_name("MX01", code), t, round(dr - cr, 2), "MXN"]
           for p, code, t, dr, cr in tb_lines("MX01", L["MX01"])]
sheet("MX01 Ledger (post-acq)", ["Entity Code", "Period", "Account Code", "Account Name", "Account Type", "Amount",
                                 "Currency"], mx_rows)

# ---------------------------------------------------------------- 02 cash & bank
cash_rows = []
BANKS = {"HQ01": [("JPMorgan Chase - Operating", "4471", "USD"), ("JPMorgan Chase - LC Collateral", "4472", "USD"),
                  ("Deutsche Bank - EUR Account", "8810", "EUR")],
         "DE01": [("Commerzbank - Geschaeftskonto", "2290", "EUR")], "BR01": [("Itau - Conta Corrente", "5531", "BRL")],
         "TR01": [("Garanti BBVA - Vadesiz", "7718", "TRY")], "MX01": [("BBVA Mexico - Operativa", "6604", "MXN")],
         "UK01": [("Barclays - Operating", "3127", "GBP")], "HLD1": [("Butterfield - Current", "0912", "USD")]}
OPEN_CASH = {"HQ01": 32_000_000, "DE01": 4_000_000, "BR01": 9_000_000, "TR01": 90_000_000, "MX01": MX_OPEN["1000"],
             "UK01": 2_500_000, "HLD1": 50_000}
for e, accts in BANKS.items():
    prev = None
    for r in L[e]:
        p, cash = r["period"], r["bs"]["1000"]
        if e == "HQ01":
            eur_usd = HQ_EUR_ACCOUNT * SPOT["EUR"][p]
            ends = [cash - HQ_RESTRICTED - eur_usd, HQ_RESTRICTED, HQ_EUR_ACCOUNT]
            begs = prev or [OPEN_CASH["HQ01"] - HQ_RESTRICTED - HQ_EUR_ACCOUNT * SPOT["EUR"]["2024-12"],
                            HQ_RESTRICTED, HQ_EUR_ACCOUNT]
        else:
            ends, begs = [cash], prev or [OPEN_CASH[e]]
        for (name, last4, ccy), b, en in zip(accts, begs, ends):
            note = ""
            if "Collateral" in name:
                note = "Restricted - cash collateral for letters of credit"
            if e == "TR01":
                note = "Repatriation subject to central bank approval (capital controls)"
            cash_rows.append([e, p, name, last4, ccy, round(b, 2), round(en, 2), "Y" if "Collateral" in name else "N",
                              note])
        prev = ends
CASHH = ["Entity Code", "Period", "Bank/Account Name", "Account Number (last 4)", "Currency", "Beginning Balance",
         "Ending Balance", "Restricted Cash", "Notes"]
sheet("02_CashBank_Position", CASHH, cash_rows)
hidden = sheet("Copy of 02_CashBank_Position", CASHH, [r for r in cash_rows if r[0] == "HQ01"])
hidden.sheet_state = "hidden"

# ---------------------------------------------------------------- 03 headcount
def fmt_date(d, f):
    if d is None:
        return None
    return d if f == "ISO" else d.strftime(FMT[f])


hc_rows = [[r["e"], r["id"], r["dept"], r["title"], "Terminated" if r["term"] else "Active", fmt_date(r["hire"], r["fmt"]),
            fmt_date(r["term"], r["fmt"]), r["ttype"], r["fte"], r["cost"], r["ccy"], r["country"]] for r in ROSTER]
sheet("03_Headcount", ["Entity Code", "Employee ID", "Department", "Job Title", "Employment Status", "Hire Date",
                       "Termination Date", "Termination Type", "FTE %", "Monthly Fully-Loaded Cost", "Currency",
                       "Country/Location"], hc_rows)

# ---------------------------------------------------------------- 04 bookings
CUSTOMERS = {"HQ01": ["Atlas Automotive", "Brightline Rail", "Corvane Energy", "Delmar Packaging", "Eastgate Foods",
                      "Fulton Aerospace", "Granite Ridge Mining", "Harbor Point Marine", "Ironwood Tools",
                      "Juniper Medical", "Keystone Logistics", "Lakeshore Chemicals",
                      "Halcyon Family Trust Ltd (related party)"],
             "DE01": ["Nordwerk GmbH", "Alpenstahl AG", "Rheinmetall Service KG", "Bavaria Pack GmbH",
                      "Elbe Logistik GmbH", "Schwarzwald Technik GmbH"],
             "MX01": ["Grupo Sierra Madre", "Industrias del Bajio", "Transportes Pacifico", "Aceros del Norte"],
             "UK01": ["Thames Valley Engineering", "Pennine Plastics", "Clyde Marine Services"]}
PRODUCTS = {"HQ01": ["Industrial Controls", "Field Services"], "DE01": ["Distribution", "Field Services"],
            "MX01": ["Industrial Controls", "Field Services"], "UK01": ["Industrial Controls", "Field Services"]}
REGION = {"HQ01": "North America", "DE01": "EMEA", "MX01": "LATAM", "UK01": "EMEA"}
bk_rows = []
deal = 9000


def booking(e, d, cust, rtype, amt, months_term=12):
    global deal
    deal += 1
    f = {"HQ01": "MDY", "DE01": "DOT", "MX01": "DMY", "UK01": "DMY"}[e]
    start = date(d.year + (d.month // 12), d.month % 12 + 1, 1)
    end = date(start.year + (start.month + months_term - 2) // 12, (start.month + months_term - 2) % 12 + 1, 28)
    bk_rows.append([e, d.strftime(FMT[f]), cust, f"D-{deal}", rtype, random.choice(PRODUCTS[e]), REGION[e],
                    round(amt, 2), CCY[e], start.strftime(FMT[f]), end.strftime(FMT[f]),
                    "Y" if rtype in ("New Business", "Expansion", "Renewal") else "N"])


for e, n, base, months in (("HQ01", 110, 260_000, range(1, 13)), ("DE01", 50, 140_000, range(1, 13)),
                           ("MX01", 25, 2_600_000, range(7, 13)), ("UK01", 20, 90_000, range(1, 10))):
    for _ in range(n):
        m = random.choice(list(months))
        rtype = random.choices(["New Business", "Expansion", "Renewal", "Contraction"], [40, 25, 30, 5])[0]
        amt = base * random.uniform(0.4, 1.8) * (-0.3 if rtype == "Contraction" else 1)
        booking(e, date(FY, m, random.randint(1, 28)), random.choice(CUSTOMERS[e][1:]), rtype, amt)
for m in (2, 5, 8, 10, 12):  # the concentration
    booking("HQ01", date(FY, m, random.randint(3, 26)), "Atlas Automotive", "Expansion" if m > 2 else "Renewal",
            random.uniform(1_400_000, 1_900_000))
booking("DE01", date(FY, 8, 14), "Nordwerk GmbH", "Churn", -900_000)
booking("HQ01", date(2026, 1, 15), "Brightline Rail", "Renewal", 420_000)  # future-dated
sheet("04_Revenue_Bookings", ["Entity Code", "Period/Transaction Date", "Customer ID/Name", "Deal/Contract ID",
                              "Revenue Type", "Product/Segment", "Region", "Amount", "Currency", "Contract Start Date",
                              "Contract End Date", "Recurring"], bk_rows)

# ---------------------------------------------------------------- 05 budget
bud = []


def b(e, p, code, amt, ccy, version="Original Budget"):
    bud.append([e, f"FY{FY}", p, code, acct_name(e, code), TYPE[code], round(amt, 2), ccy, version])


for i, p in enumerate(M):
    s = {a: SAL[("HQ01", a)][p] for a in ("5100", "6000", "6200", "6300")}
    prod = 3_250_000 * 1.006 ** i
    for code, amt in (("4000", prod), ("4100", 820_000), ("5000", 0.54 * (prod + HQ_IC[p])), ("5100", s["5100"]),
                      ("6000", s["6000"] * 0.85), ("6100", 140_000), ("6200", s["6200"]), ("6300", s["6300"] * 0.98),
                      ("6400", 120_000), ("6600", 80_000)):
        b("HQ01", p, code, amt, "USD")
        if MNUM[p] >= 7 and code in ("4000", "4100"):
            b("HQ01", p, code, amt * (0.98 if code == "4000" else 0.97), "USD", "Reforecast (Jul)")
    de_prod = 1_950_000 * 1.003 ** i
    for code, amt in (("4000", de_prod), ("4100", 260_000), ("5000", 0.48 * de_prod), ("5900", 520_000),
                      ("6000", SAL[("DE01", "6000")][p] * 1.21), ("6100", 55_000),
                      ("6300", SAL[("DE01", "6300")][p] * 1.21), ("6400", 70_000)):
        b("DE01", p, code, amt * 1.05, "USD", "Original Budget (USD at budget rate 1.05)")
    for code, amt in (("4000", 5_300_000 * 1.004 ** i), ("4100", 820_000), ("5000", 0.60 * 5_300_000 * 1.004 ** i),
                      ("6000", SAL[("BR01", "6000")][p]), ("6300", SAL[("BR01", "6300")][p]), ("6400", 150_000)):
        b("BR01", p, code, amt * random.uniform(0.98, 1.02), "BRL")
    for code, amt in (("4000", 960_000), ("4100", 150_000), ("5000", 557_000), ("6000", SAL[("UK01", "6000")][M[0]]),
                      ("6300", SAL[("UK01", "6300")][M[0]])):
        b("UK01", p, code, amt, "GBP")  # full-year budget, including Q4 after the disposal
    for code, amt in (("4000", 32_000_000), ("4100", 6_000_000), ("5000", 25_600_000), ("6400", 1_500_000)):
        b("TR01", p, code, amt, "TRY")  # nominal, set before the year's inflation
ref = next(r for r in bud if r[0] == "HQ01" and r[2] == "2025-10" and r[3] == "6300")
ref[6] = "#REF!"
sheet("05_Budget_Forecast", ["Entity Code", "Fiscal Year", "Period", "Account Code", "Account Name", "Account Type",
                             "Budgeted Amount", "Currency", "Version"], bud)

# ---------------------------------------------------------------- 06 FX
fx_rows = []
for c in ("EUR", "GBP", "BRL", "TRY", "MXN", "INR", "CAD"):
    for p in M:
        fx_rows.append([p, c, "USD", "Period-End Spot", SPOT[c][p], "Group Treasury"])
        if not (c == "MXN" and p == "2025-08"):  # missing average rate
            fx_rows.append([p, c, "USD", "Period-Average", AVG[c][p], "Group Treasury"])
sheet("06_FX_Rates", ["Period", "From Currency", "To Currency", "Rate Type", "Rate", "Source"], fx_rows)

# ---------------------------------------------------------------- 07 intercompany
ic_rows = []
for p in M:
    ic_rows.append([p, "HQ01", "DE01", "Sale of goods", round(HQ_IC[p], 2), "USD", "Finished goods for EU distribution"])
    ic_rows.append([p, "HQ01", "BR01", "Management fee", 50_000, "USD", "Group services charge"])
sheet("07_Intercompany_Detail", ["Period", "Entity (From)", "Entity (To)", "Transaction Type", "Amount", "Currency",
                                 "Description"], ic_rows)

# ---------------------------------------------------------------- 08 debt & covenants (company's own calculation)
def ltm_like(qe):  # the company annualizes year-to-date EBITDA and EXCLUDES the revolver from net debt
    ps = M[:M.index(qe) + 1]
    e_ann = group_ebitda(L, ps) * 12 / len(ps)
    i_ann = interest(L, ps) * 12 / len(ps)
    nd = term_loan - group_cash(L, qe)
    return nd / e_ann, e_ann / i_ann


debt_rows = []
NL_DEF = ("Net debt / LTM Adjusted EBITDA. Net debt = term loan + revolver - unrestricted cash. Adjusted EBITDA "
          "adds back restructuring, acquisition costs and other one-time items.")
IC_DEF = "LTM Adjusted EBITDA / LTM interest expense"
for qe in ("2025-03", "2025-06", "2025-09", "2025-12"):
    nl, icov = ltm_like(qe)
    rev_drawn = L["HQ01"][M.index(qe)]["bs"]["2400"]
    for fac, ftype, lim, drawn, mat, rate in (("Term Loan A", "Term loan", term_loan, term_loan, "2026-06-30", 0.0525),
                                              ("Revolving Credit Facility", "Revolver", 10_000_000, rev_drawn,
                                               "2027-03-31", 0.07)):
        for cname, cdef, thr, cmp_, val in (("Net leverage", NL_DEF, COVENANT_LEVERAGE, "<=", nl),
                                            ("Interest cover", IC_DEF, 4.0, ">=", icov)):
            debt_rows.append(["HQ01", qe, fac, ftype, lim, drawn, "USD", mat, rate, cname, cdef, thr, cmp_,
                              round(val, 2)])
sheet("08_Debt_Covenants", ["Entity Code", "Period", "Facility Name", "Facility Type", "Limit", "Drawn", "Currency",
                            "Maturity Date", "Interest Rate", "Covenant Name", "Covenant Definition", "Threshold",
                            "Comparison", "Reported Value"], debt_rows)

# ---------------------------------------------------------------- 09 operational KPIs (months laid out as columns)
mcols = [date(FY, MNUM[p], 1).strftime("%b-%y") for p in M]
kpi = [["HQ01", "Units shipped", "units"] + [round(L["HQ01"][i]["pnl"]["4000"] / 245) for i in range(12)],
       ["HQ01", "Capacity utilization", "%"] + [round(0.78 + 0.007 * i, 3) for i in range(12)],
       ["DE01", "Orders processed", "count"] + [round(L["DE01"][i]["pnl"]["4000"] / 1_900) for i in range(12)]]
sheet("09_Operational_KPIs", ["Entity Code", "Metric", "Unit"] + mcols, kpi)

# ---------------------------------------------------------------- 10 adjustments & events (DE01 release NOT listed)
ev = [["HQ01", "2025-07", "Acquisition", "Acquired 100% of MX01 on 1 Jul 2025. Consideration USD 25.0m: 22.0m cash "
       "plus a 3.0m earn-out at fair value. Purchase price allocation not yet performed.", MX_CONSIDERATION, "USD", "Y"],
      ["HQ01", "2025-07", "Other", "Acquisition advisory fees (MX01)", 650_000, "USD", "Y"],
      ["HQ01", "2025-09", "Disposal", "Sold 100% of UK01 (the group's only UK operation) on 30 Sep 2025 for USD 9.0m "
       "cash.", UK_PROCEEDS, "USD", "Y"],
      ["HQ01", "2025-11", "Other", "Drew USD 5.0m on the revolving credit facility.", 5_000_000, "USD", "N"]]
ev += [["HQ01", p, "Restructuring", "Plant footprint optimisation", 120_000, "USD", "Y"]
       for p in M if MNUM[p] in RESTRUCT_MONTHS]
ev += [["DE01", "2025-03", "Restructuring", "Sales office consolidation", 400_000, "EUR", "Y"]]
sheet("10_Adjustments_Events", ["Entity Code", "Period", "Event Type", "Description", "Amount", "Currency",
                                "One-Time"], ev)

wb.save(OUT_XLSX)

# ---------------------------------------------------------------- expected values (the answer key's numbers)
def fy_ni(e, ccy, periods=None):
    return sum(r["ni"] * AVG[ccy][r["period"]] for r in L[e] if periods is None or r["period"] in periods)


mx_net_assets = sum(v for k, v in MX_OPEN.items() if TYPE[k] == "Asset") - \
    sum(v for k, v in MX_OPEN.items() if TYPE[k] == "Liability")
mx_na_usd = mx_net_assets * SPOT["MXN"]["2025-06"]
uk_sep = L["UK01"][-1]
uk_na = sum(v for k, v in uk_sep["bs"].items() if TYPE[k] == "Asset") - \
    sum(v for k, v in uk_sep["bs"].items() if TYPE[k] == "Liability")
uk_na_usd = uk_na * SPOT["GBP"]["2025-09"]
ebitda_fy = group_ebitda(L, M)
cash_dec = group_cash(L, M[-1])
revolver_dec = L["HQ01"][-1]["bs"]["2400"]
nd_true = term_loan + revolver_dec - cash_dec
nl_true = nd_true / ebitda_fy
hq = {r["period"]: r for r in L["HQ01"]}
de = {r["period"]: r for r in L["DE01"]}
gross = defaultdict(float)
for r in bk_rows:
    if r[7] > 0 and r[1][-4:] == str(FY):
        gross[r[2]] += r[7] * (AVG[CCY[r[0]]]["2025-06"])
expected = {
    "description": "Reference values for a correct FY2025 cycle on Business_Review_Complex_Input_FY2025.xlsx. "
                   "USD figures use the file's own rates (average for P&L, period-end spot for balances).",
    "term_loan_usd": term_loan,
    "covenants": {
        "group_adjusted_ebitda_fy_usd": round(ebitda_fy), "unrestricted_cash_dec_usd": round(cash_dec),
        "revolver_dec_usd": revolver_dec, "net_debt_incl_revolver_usd": round(nd_true),
        "net_leverage_recomputed": round(nl_true, 2), "net_leverage_threshold": COVENANT_LEVERAGE,
        "net_leverage_headroom_pct": round((COVENANT_LEVERAGE - nl_true) / COVENANT_LEVERAGE * 100, 1),
        "net_leverage_company_reported_q4": round(ltm_like("2025-12")[0], 2),
        "why_company_figure_differs": "Company excludes the revolver from net debt",
        "interest_cover_recomputed": round(ebitda_fy / interest(L, M), 2),
        "notes": "EBITDA here is as-reported for the year (UK01 Jan-Sep, MX01 Jul-Dec), translated at monthly "
                 "average rates, excluding intercompany lines; a pro-forma covenant view would differ.",
    },
    "mx01_acquisition": {"consideration_usd": MX_CONSIDERATION, "cash_paid_usd": MX_CASH,
                         "earnout_fair_value_usd": MX_EARNOUT, "net_assets_at_acquisition_mxn": mx_net_assets,
                         "net_assets_at_acquisition_usd_jun_spot": round(mx_na_usd),
                         "provisional_goodwill_usd": round(MX_CONSIDERATION - mx_na_usd),
                         "deal_costs_usd": 650_000, "earnout_fv_change_dec_usd": 300_000,
                         "fy_net_income_post_acq_usd": round(fy_ni("MX01", "MXN"))},
    "uk01_disposal": {"proceeds_usd": UK_PROCEEDS, "parent_book_gain_usd": UK_PROCEEDS - UK_COST,
                      "net_assets_30sep_gbp": round(uk_na), "net_assets_30sep_usd": round(uk_na_usd),
                      "consolidated_gain_before_cta_recycling_usd": round(UK_PROCEEDS - uk_na_usd),
                      "jan_sep_net_income_usd": round(fy_ni("UK01", "GBP")),
                      "treatment": "Discontinued operation (major geographic area); restate KPIs on continuing ops"},
    "br01_nci": {"ownership_pct": 70,
                 "fy_net_income_usd_all_12_months": round(fy_ni("BR01", "BRL")),
                 "nci_share_30pct_usd_all_12_months": round(0.30 * fy_ni("BR01", "BRL")),
                 "april_net_income_missing_from_export_usd": round(
                     fy_ni("BR01", "BRL", ["2025-04"]))},
    "jv01_equity_method": {"ownership_pct": 40, "fy_net_income_inr": round(sum(r["ni"] for r in L["JV01"])),
                           "share_of_profit_usd": round(0.40 * fy_ni("JV01", "INR")),
                           "fiscal_year_start": "April"},
    "hq01": {"recurring_restructuring_usd": 120_000 * len(RESTRUCT_MONTHS),
             "restructuring_months": [f"{FY}-{m:02d}" for m in RESTRUCT_MONTHS],
             "dec_bill_and_hold_uplift_usd": round(hq["2025-12"]["pnl"]["4000"] / 1.35 * 0.35),
             "dec_dso_days": 60, "prior_dso_days": 45,
             "capitalized_dev_monthly_h1_usd": 150_000, "capitalized_dev_monthly_h2_usd": 400_000,
             "period13_adjustment_usd": 85_000, "future_dated_rows_period": "2026-02",
             "related_party_customer": "Halcyon Family Trust Ltd (related party)",
             "parent_investment_in_subsidiaries_dec_usd": hq["2025-12"]["bs"]["1500"],
             "investment_in_jv_at_cost_usd": 4_000_000},
    "de01": {"warranty_provision_release_nov_eur": 600_000, "dec_ic_in_transit_eur": 150_000,
             "gl_salary_vs_roster_factor": 1.21, "export": "Year-to-date cumulative P&L"},
    "tr01": {"spot_step_change_mar": round(SPOT["TRY"]["2025-03"] / SPOT["TRY"]["2025-02"] - 1, 3),
             "fy_net_income_try": round(sum(r["ni"] for r in L["TR01"])),
             "hyperinflation_index_supplied": False},
    "fx": {"missing_rates": ["MXN 2025-08 Period-Average"]},
    "concentration": {"top_customer": max(gross, key=gross.get),
                      "top_customer_share_of_gross_bookings_approx": round(max(gross.values()) / sum(gross.values()), 3)},
    "fy_net_income_native": {e: round(sum(r["ni"] for r in L[e])) for e in L},
}
with open(OUT_JSON, "w", encoding="utf-8") as f:
    json.dump(expected, f, indent=2)
print(f"Wrote {OUT_XLSX}\nWrote {OUT_JSON}")
print(f"Term loan {term_loan:,.0f} | group EBITDA {ebitda_fy:,.0f} | net leverage {nl_true:.2f}x "
      f"(company reports {expected['covenants']['net_leverage_company_reported_q4']}x) | "
      f"interest cover {expected['covenants']['interest_cover_recomputed']}x")
print(f"Rows: GL {len(rows)} + BR01 {len(br_rows)} + MX01 {len(mx_rows)} | headcount {len(hc_rows)} | bookings "
      f"{len(bk_rows)} | budget {len(bud)} | cash {len(cash_rows)}")
