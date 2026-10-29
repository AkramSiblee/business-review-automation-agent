#!/usr/bin/env python3
"""
Phase 1 / 2.5 helper -- deterministic first pass over whatever is in the resolved inputs folder.

  python scripts/profile_inputs.py                     # profile + write the Data Validation Report
  python scripts/profile_inputs.py --period 2026-Q1    # override the derived period label
  python scripts/profile_inputs.py --notes notes.json  # add Claude's judgment items to the report
  python scripts/profile_inputs.py --no-report         # profile only, print summary

Works on one multi-tab workbook (the standard case) or several files (tabs pooled per category).
docx/pdf/pptx inputs are listed as "needs extraction" -- Claude extracts those itself (Phase 1).

Writes to the resolved outputs folder:
  <Period>_Data_Validation_Report.xlsx   one sheet, skimmable in seconds (Phase 2.5)
  _input_profile.json                    machine-readable profile, so a resumed or fresh agent can
                                         continue Phase 3 without re-deriving the triage
Prints a short summary only -- never raw rows.

--notes JSON format: [{"tier": "LOW", "item": "...", "detail": "..."}, ...]
"""
import argparse
import json
import os
import sys
from collections import Counter, defaultdict
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import brlib  # noqa: E402

ENTITY_COLS = ["Entity Code", "Entity (From)", "Entity (To)"]
PERIOD_COLS = ["Period", "Period/Transaction Date"]
DATE_HINTS = ("date", "period")


def col(row, colmap, name):
    actual = colmap.get(name)
    return row.get(actual) if actual else None


def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def profile(inputs_dir):
    files = brlib.list_input_files(inputs_dir)
    tables, unreadable, needs_extraction = [], [], []
    deferred = []  # (table, rows, colmap, pcol): checks that need the fiscal-year start from config
    for path in files:
        ext = os.path.splitext(path)[1].lower()
        if ext in (".docx", ".pdf", ".pptx", ".xls"):
            needs_extraction.append(os.path.basename(path))
            continue
        try:
            for label, tab, headers, rows, meta in brlib.load_tables(path):
                cat, conf, cov, colmap = brlib.classify(tab, headers)
                spec = brlib.CATEGORIES.get(cat, {})
                entities = set()
                for r in rows:
                    for c in ENTITY_COLS:
                        if colmap.get(c):
                            entities.add(r.get(colmap[c]))
                pcol = next((colmap[c] for c in PERIOD_COLS if colmap.get(c)), None)
                periods = {brlib.to_period(r.get(pcol)) for r in rows} if pcol else set()
                # Per entity: one entity's ISO dates must not mask another entity's ambiguous DD/MM ones.
                date_evidence = {}
                ecol = colmap.get("Entity Code")
                for h in headers:
                    if any(k in h.lower() for k in DATE_HINTS):
                        groups = defaultdict(list)
                        for r in rows:
                            groups[str(r.get(ecol)) if ecol else "all"].append(r.get(h))
                        for ent, vals in sorted(groups.items()):
                            ev = brlib.date_order_evidence(vals)
                            if ev:
                                date_evidence[h if ent == "all" else f"{h} [{ent}]"] = ev
                dupes = len(rows) - len({tuple(sorted((k, str(v)) for k, v in r.items())) for r in rows})
                t = {
                    "source": label, "tab": tab, "category": cat, "confidence": conf,
                    "header_coverage": round(cov, 2), "rows": len(rows),
                    "missing_columns": [c for c in spec.get("columns", []) if c not in colmap],
                    "extra_columns": [h for h in headers if h not in colmap.values()],
                    "entities": sorted(str(e) for e in entities if e not in (None, "")),
                    "periods": sorted(p for p in periods if p),
                    "unparsed_periods": sum(1 for r in rows if pcol and r.get(pcol) not in (None, "")
                                            and not brlib.to_period(r.get(pcol))
                                            and not brlib.ADJ_PERIOD_RE.search(str(r.get(pcol)))),
                    "currencies": sorted({str(r.get(colmap["Currency"])) for r in rows
                                          if colmap.get("Currency") and r.get(colmap["Currency"])}),
                    "date_order_evidence": date_evidence,
                    "exact_duplicate_rows": dupes,
                }
                if cat == "01_GL_TrialBalance" and colmap.get("Debit") and colmap.get("Credit"):
                    t["tb_imbalances"] = tb_imbalances(rows, colmap, pcol)
                if cat == "00_Company_Config":
                    t["config_rows"] = [{k: (str(v) if v is not None else None) for k, v in r.items()}
                                        for r in rows]
                t["shape_checks"] = shape_checks(t, cat, headers, rows, colmap, pcol, meta)
                if cat == "01_GL_TrialBalance":
                    deferred.append((t, rows, colmap, pcol))
                tables.append(t)
        except Exception as e:  # corrupt / password-protected / still syncing
            unreadable.append({"file": os.path.basename(path), "error": f"{type(e).__name__}: {e}"})
    fy_start = config_fy_start([r for t in tables for r in t.get("config_rows", [])])
    for t, rows, colmap, pcol in deferred:
        t["shape_checks"] += gl_checks(t, rows, colmap, pcol, fy_start)
    return {"inputs_dir": inputs_dir, "files": [os.path.basename(f) for f in files], "tables": tables,
            "needs_extraction": needs_extraction, "unreadable": unreadable}


def chk(check, scope, result, detail):
    return {"check": check, "scope": scope, "result": result, "detail": detail}


def config_fy_start(config_rows):
    for r in config_rows:
        v = str(r.get("Fiscal Year Start Month") or "").strip()
        names = ["january", "february", "march", "april", "may", "june", "july", "august", "september",
                 "october", "november", "december"]
        for i, name in enumerate(names, start=1):  # "January", "Apr", "Sept"
            if len(v) >= 3 and name.startswith(v.lower().rstrip(".")):
                return i
        try:
            return int(float(v))
        except ValueError:
            pass
    return 1


def amount(row, colmap, headers):
    """Signed debit-positive amount: Debit - Credit, or a single Amount/Balance/Value column."""
    if colmap.get("Debit") or colmap.get("Credit"):
        return num(col(row, colmap, "Debit")) - num(col(row, colmap, "Credit"))
    h = signed_amount_col(colmap, headers)
    return num(row.get(h)) if h else 0.0


def signed_amount_col(colmap, headers):
    if colmap.get("Debit") or colmap.get("Credit"):
        return None
    return next((h for h in headers if h.strip().lower() in brlib.SIGNED_AMOUNT_HEADERS), None)


def shape_checks(t, cat, headers, rows, colmap, pcol, meta):
    """Data-shape traps from complex-scenarios-playbook.md section 2 that apply to any tab."""
    out, src = [], t["source"]
    if meta["hidden"] or brlib.SCRATCH_TAB_RE.search(t["tab"]):
        out.append(chk("Hidden / scratch tab", src, "REVIEW",
                       ("Hidden sheet" if meta["hidden"] else "Tab name looks like scratch work") +
                       " -- exclude unless it is the only source for its category (playbook 2.11)"))
    period_headers = [h for h in headers if brlib.to_period(h)]
    if len(period_headers) >= 3:
        out.append(chk("Months laid out as columns", src, "REVIEW",
                       f"{len(period_headers)} period headers ({period_headers[0]} .. {period_headers[-1]}) -- "
                       f"unpivot to one row per period before use (playbook 2.3)"))
    scale = brlib.UNIT_SCALE_RE.search(" | ".join([t["tab"], meta["preamble"]] + headers))
    if scale:
        out.append(chk("Unit scale stated", src, "REVIEW",
                       f"'{scale.group(0)}' found in title/headers -- scale to actual units (playbook 2.5)"))
    errors = Counter(v.strip() for r in rows for v in r.values()
                     if isinstance(v, str) and v.strip() in brlib.EXCEL_ERRORS)
    if errors:
        out.append(chk("Excel error values", src, "FAIL",
                       f"{sum(errors.values())} cells ({', '.join(f'{k} x{n}' for k, n in errors.items())}) -- "
                       f"treat as missing, never zero; source workbook broken upstream (playbook 2.6)"))
    if pcol:
        adj = sum(1 for r in rows if isinstance(r.get(pcol), str) and brlib.ADJ_PERIOD_RE.search(r.get(pcol)))
        if adj:
            out.append(chk("Adjustment period (13+/ADJ)", src, "REVIEW",
                           f"{adj} rows in an adjustment period -- fold into the last fiscal month and log "
                           f"(playbook 2.10)"))
        today = f"{date.today():%Y-%m}"
        future = sorted(p for p in t["periods"] if p > today)
        if future and cat in ("01_GL_TrialBalance", "02_CashBank_Position", "03_Headcount"):
            out.append(chk("Future-dated actuals", src, "REVIEW",
                           f"Periods after {today}: {', '.join(future)} -- likely forecast leakage (playbook 2.12)"))
    ecol, ccol = colmap.get("Entity Code"), colmap.get("Currency")
    if ecol and ccol and cat in ("01_GL_TrialBalance", "03_Headcount", "04_Revenue_Bookings", "05_Budget_Forecast"):
        per_ent = defaultdict(set)
        for r in rows:
            if r.get(ccol):
                per_ent[str(r.get(ecol))].add(str(r.get(ccol)))
        for ent, ccys in sorted(per_ent.items()):
            if len(ccys) > 1:
                out.append(chk("Mixed currencies within an entity", f"{src} / {ent}", "REVIEW",
                               f"{', '.join(sorted(ccys))} -- translate each row to functional currency at the "
                               f"same-period rate (playbook 2.8)"))
    if cat in ("01_GL_TrialBalance", "05_Budget_Forecast") and colmap.get("Account Name"):
        acode = colmap.get("Account Code")
        subt = [r for r in rows if brlib.SUBTOTAL_RE.search(str(r.get(colmap["Account Name"]) or ""))
                or (acode and r.get(acode) in (None, "") and amount(r, colmap, headers))]
        if subt:
            names = brlib.counter_summary(str(r.get(colmap["Account Name"])) for r in subt)
            out.append(chk("Subtotal / total rows", src, "REVIEW",
                           f"{len(subt)} rows look like subtotals ({names}) -- drop before summing, "
                           f"then re-check debits = credits (playbook 2.2)"))
    if cat == "06_FX_Rates":
        out += fx_checks(src, rows, colmap)
    return out


def gl_checks(t, rows, colmap, pcol, fy_start):
    """GL-specific traps: signed amount column, YTD-cumulative balances, missing months, CoA changes."""
    out, src = [], t["source"]
    headers = list(rows[0].keys()) if rows else []
    signed = signed_amount_col(colmap, headers)
    tcol, ecol, acol = colmap.get("Account Type"), colmap.get("Entity Code"), colmap.get("Account Code")
    if signed:
        rev = [num(r.get(signed)) for r in rows if tcol and "revenue" in str(r.get(tcol)).lower()]
        conv = ("credit-negative (ERP convention)" if rev and sum(v < 0 for v in rev) > len(rev) / 2
                else "debit-negative / revenue positive" if rev else "unknown -- no Account Type to infer from")
        out.append(chk("Single signed amount column", src, "REVIEW",
                       f"'{signed}' instead of Debit/Credit; sign convention looks {conv} -- normalize, then "
                       f"confirm assets = liabilities + equity (playbook 2.4)"))
    if not pcol:
        return out
    by_ent_period = defaultdict(set)
    series = defaultdict(dict)
    for r in rows:
        p = brlib.to_period(r.get(pcol))
        if not p:
            continue
        ent = str(r.get(ecol)) if ecol else "all"
        acct = str(r.get(acol)) if acol else str(r.get(colmap.get("Account Name")))
        by_ent_period[ent].add(p)
        if tcol and any(k in str(r.get(tcol)).lower() for k in brlib.PNL_TYPE_HINTS) and \
                not any(k in str(r.get(tcol)).lower() for k in brlib.BS_TYPE_HINTS):
            key = (ent, acct, brlib.fiscal_year(p, fy_start))
            series[key][p] = series[key].get(p, 0.0) + amount(r, colmap, headers)
    today = f"{date.today():%Y-%m}"
    for ent, ps in sorted(by_ent_period.items()):
        closed = {p for p in ps if p <= today}  # future-dated rows are reported separately
        holes = sorted(set(brlib.month_range(min(closed), max(closed))) - closed) if closed else []
        if holes:
            out.append(chk("Missing months in GL", f"{src} / {ent}", "REVIEW",
                           f"No rows for {', '.join(holes)} -- never interpolate; mark comparisons n/a "
                           f"(playbook 2.7)"))

    # YTD-cumulative: most P&L accounts only ever grow in absolute size within a fiscal year.
    # Tested per entity -- one entity exporting YTD must not be diluted by others exporting monthly.
    def cumulative(s):
        vals = [abs(s[p]) for p in sorted(s)]
        return all(b >= a - 0.005 for a, b in zip(vals, vals[1:])) and vals[-1] > 1.5 * max(vals[0], 0.01)

    per_ent = defaultdict(list)
    for (ent, _, _), s in series.items():
        if len(s) >= 4:
            per_ent[ent].append(cumulative(s))
    for ent, flags in sorted(per_ent.items()):
        share = sum(flags) / len(flags)
        if share > 0.6:
            out.append(chk("P&L looks YTD-cumulative", f"{src} / {ent}", "REVIEW",
                           f"{share:.0%} of {ent}'s P&L account series rise every month within the fiscal year -- "
                           f"difference to monthly activity, then re-run the RE rollforward (playbook 2.1)"))
    if not tcol:
        out.append(chk("P&L looks YTD-cumulative", src, "REVIEW",
                       "No Account Type column, so cumulative vs monthly could not be tested -- check manually"))
    # Chart-of-accounts continuity per entity.
    if acol and ecol:
        accts = defaultdict(lambda: defaultdict(set))
        for r in rows:
            p = brlib.to_period(r.get(pcol))
            if p:
                accts[str(r.get(ecol))][p].add(str(r.get(acol)))
        for ent, per in sorted(accts.items()):
            union = set().union(*per.values())
            partial = sorted(a for a in union if sum(a in s for s in per.values()) < len(per))
            if partial and len(per) > 1:
                out.append(chk("Account set changes across periods", f"{src} / {ent}", "INFO",
                               f"{len(partial)} accounts not present every period "
                               f"({brlib.counter_summary(partial, 6)}) -- check for remaps vs genuine "
                               f"discontinuations (playbook 2.9)"))
    return out


def fx_checks(src, rows, colmap):
    """Missing spot/average pairs and >10% monthly spot moves (devaluation / hyperinflation watch)."""
    out = []
    need = ("Period", "From Currency", "To Currency", "Rate Type", "Rate")
    if not all(colmap.get(c) for c in need):
        return out
    rates = defaultdict(dict)
    for r in rows:
        p = brlib.to_period(col(r, colmap, "Period"))
        rt = str(col(r, colmap, "Rate Type") or "").lower()
        kind = "spot" if any(k in rt for k in ("spot", "end", "clos")) else "avg" if "av" in rt else None
        if p and kind:
            rates[(str(col(r, colmap, "From Currency")), str(col(r, colmap, "To Currency")), kind)][p] = \
                num(col(r, colmap, "Rate"))
    pairs = {(f, to) for f, to, _ in rates}
    for f, to in sorted(pairs):
        spot, avg = rates.get((f, to, "spot"), {}), rates.get((f, to, "avg"), {})
        one_sided = sorted(set(spot) ^ set(avg))
        if one_sided:
            out.append(chk("FX spot/average pair incomplete", f"{src} / {f}->{to}", "REVIEW",
                           f"Only one rate type for {', '.join(one_sided)} -- missing average: use spot as proxy "
                           f"(MEDIUM); missing closing spot: LOW (playbook 4.3)"))
        ps = sorted(spot)
        for a, b in zip(ps, ps[1:]):
            if spot[a] and abs(spot[b] / spot[a] - 1) > 0.10:
                out.append(chk("FX step change > 10%", f"{src} / {f}->{to}", "REVIEW",
                               f"Spot {a} {spot[a]:g} -> {b} {spot[b]:g} ({spot[b] / spot[a] - 1:+.1%}) -- "
                               f"devaluation / hyperinflation watch; split variance at constant currency "
                               f"(playbook 4.1-4.2)"))
    return out


def config_checks(config_rows):
    """Group-structure flags from 00_Company_Config (playbook section 3)."""
    out = []

    def pct(v):
        try:
            x = float(str(v).replace("%", "").strip())
        except (TypeError, ValueError):
            return None
        return x * 100 if 0 < x <= 1 and "%" not in str(v) else x

    for r in config_rows:
        ent = r.get("Entity Code") or "?"
        own, method = pct(r.get("Ownership %")), str(r.get("Consolidation Method") or "").strip().lower()
        if method in ("equity", "equity method", "associate", "jv", "joint venture"):
            out.append(chk("Equity-method entity", ent, "INFO",
                           f"Method '{r.get('Consolidation Method')}' ({own:g}% owned) -- one-line share of profit, "
                           f"excluded from group revenue/opex/headcount/cash (playbook 3.2)" if own is not None
                           else f"Method '{r.get('Consolidation Method')}' -- one-line share of profit (playbook 3.2)"))
        elif method.startswith("prop"):
            out.append(chk("Proportionate consolidation", ent, "INFO",
                           f"Include {own:g}% of each line (playbook 3.3)" if own is not None
                           else "Ownership % missing -- needed for proportionate consolidation"))
        elif own is not None and own < 100:
            if not method and 20 <= own <= 50:
                out.append(chk("Consolidation method unclear", ent, "REVIEW",
                               f"{own:g}% owned, no method stated -- full vs equity method changes revenue "
                               f"materially (LOW, playbook 3.2)"))
            else:
                out.append(chk("Non-controlling interest", ent, "INFO",
                               f"{own:g}% owned, consolidated in full -- show NCI share of net income and equity "
                               f"(playbook 3.1)"))
    fy = {str(r.get("Fiscal Year Start Month")) for r in config_rows if r.get("Fiscal Year Start Month")}
    if len(fy) > 1:
        out.append(chk("Fiscal year-ends differ", "00_Company_Config", "REVIEW",
                       f"Fiscal year start months {', '.join(sorted(fy))} -- align to parent calendar, disclose "
                       f"any lag (playbook 3.6)"))
    stds = {str(r.get("Accounting Standard")) for r in config_rows if r.get("Accounting Standard")}
    if len(stds) > 1:
        out.append(chk("Mixed accounting standards", "00_Company_Config", "INFO",
                       f"{', '.join(sorted(stds))} -- flag lease, R&D, inventory and revenue-timing divergence "
                       f"(SKILL Phase 4, playbook 8)"))
    return out


def tb_imbalances(rows, colmap, pcol):
    """Debits should equal credits per entity/period in a trial balance. Returns the failing groups."""
    sums = defaultdict(float)
    for r in rows:
        key = (str(col(r, colmap, "Entity Code")), brlib.to_period(r.get(pcol)) if pcol else None)
        sums[key] += num(col(r, colmap, "Debit")) - num(col(r, colmap, "Credit"))
    return [{"entity": e, "period": p, "debit_minus_credit": round(v, 2)}
            for (e, p), v in sorted(sums.items(), key=lambda kv: str(kv[0])) if abs(v) >= 1.0]


def analyse(prof):
    by_cat = defaultdict(list)
    for t in prof["tables"]:
        by_cat[t["category"]].append(t)

    def union(cat, field):
        return sorted({x for t in by_cat.get(cat, []) for x in t[field]})

    config_rows = [r for t in by_cat.get("00_Company_Config", []) for r in t.get("config_rows", [])]
    config_entities = sorted({r.get("Entity Code") for r in config_rows if r.get("Entity Code")})
    all_entities = config_entities or sorted({e for t in prof["tables"] for e in t["entities"]})
    all_currencies = sorted({c for t in prof["tables"] for c in t["currencies"]})
    multi = {"multi_entity": len(all_entities) > 1, "multi_currency": len(all_currencies) > 1}

    expected = []
    for cat, spec in brlib.CATEGORIES.items():
        hits = by_cat.get(cat, [])
        if hits:
            status = "Present"
        elif spec.get("optional"):
            status = "Not provided (optional)"
        elif spec.get("only_if") and not multi[spec["only_if"]]:
            status = "Not needed"
        else:
            status = "Missing"
        expected.append({"category": cat, "status": status,
                         "matched": "; ".join(f'{t["source"]} ({t["confidence"]})' for t in hits)})

    gl_periods = [p for p in union("01_GL_TrialBalance", "periods") if p <= f"{date.today():%Y-%m}"]
    # Equity-method / non-consolidated entities only need a GL (for share of profit) -- no cash, budget,
    # headcount or bookings are expected for them (playbook 3.2).
    one_line = {r.get("Entity Code") for r in config_rows
                if str(r.get("Consolidation Method") or "").strip().lower()
                in ("equity", "equity method", "associate", "jv", "joint venture", "fair value", "not consolidated")}
    gaps = []
    for cat in ("01_GL_TrialBalance", "02_CashBank_Position", "03_Headcount", "04_Revenue_Bookings",
                "05_Budget_Forecast"):
        if cat not in by_cat:
            continue
        have = set(union(cat, "entities"))
        for e in all_entities:
            if e not in have and not (e in one_line and cat != "01_GL_TrialBalance"):
                gaps.append({"type": "Entity missing", "category": cat, "detail": f"{e} has no rows in {cat}"})
    for cat in ("02_CashBank_Position", "05_Budget_Forecast"):
        if cat in by_cat and gl_periods:
            missing = sorted(set(gl_periods) - set(union(cat, "periods")))
            if missing:
                gaps.append({"type": "Period missing", "category": cat,
                             "detail": f"GL has periods not in {cat}: {', '.join(missing)}"})
    if "07_Intercompany_Detail" in by_cat or multi["multi_entity"]:
        if multi["multi_entity"] and "07_Intercompany_Detail" not in by_cat:
            gaps.append({"type": "Elimination gap", "category": "07_Intercompany_Detail",
                         "detail": "Multiple entities but no intercompany data -- elimination cannot fully run; "
                                   "must be disclosed in deliverables"})
    if multi["multi_currency"] and "06_FX_Rates" in by_cat:
        fx_periods = set(union("06_FX_Rates", "periods"))
        missing = sorted(set(gl_periods) - fx_periods)
        if missing:
            gaps.append({"type": "Period missing", "category": "06_FX_Rates",
                         "detail": f"No FX rates for GL periods: {', '.join(missing)}"})

    checks = []
    for t in prof["tables"]:
        for imb in t.get("tb_imbalances", []):
            checks.append({"check": "Trial balance debits = credits", "scope": f'{imb["entity"]} {imb["period"]}',
                           "result": "FAIL",
                           "detail": f'Debits - credits = {imb["debit_minus_credit"]:,.2f} (an unclosed-earnings '
                                     f'imbalance on a pre-close TB is expected -- confirm before flagging)'})
        if t["exact_duplicate_rows"]:
            checks.append({"check": "Exact duplicate rows", "scope": t["source"], "result": "REVIEW",
                           "detail": f'{t["exact_duplicate_rows"]} identical rows'})
        if t["unparsed_periods"]:
            checks.append({"check": "Period values readable", "scope": t["source"], "result": "REVIEW",
                           "detail": f'{t["unparsed_periods"]} period values could not be parsed'})
        for h, ev in t["date_order_evidence"].items():
            if ev in ("DD/MM", "MM/DD"):
                checks.append({"check": "Date component order", "scope": f'{t["source"]} / {h}',
                               "result": "INFO", "detail": f"{ev} proven by a component > 12 -- parse as {ev}/YYYY"})
            else:
                checks.append({"check": "Date component order", "scope": f'{t["source"]} / {h}',
                               "result": "REVIEW",
                               "detail": f"{ev}: no value proves DD/MM vs MM/DD" if ev == "AMBIGUOUS"
                               else f"{ev}: column mixes DD/MM and MM/DD values"})
    gl_tables = [t for t in by_cat.get("01_GL_TrialBalance", []) if "tb_imbalances" in t]
    if gl_tables and not any(t["tb_imbalances"] for t in gl_tables):
        checks.append({"check": "Trial balance debits = credits", "scope": "all entities/periods",
                       "result": "PASS", "detail": ""})
    # Data-shape and group-structure traps (complex-scenarios-playbook.md, marked ⚙ there).
    for t in prof["tables"]:
        checks.extend(t.get("shape_checks", []))
    checks.extend(config_checks(config_rows))

    fy_start = config_fy_start(config_rows)
    return {"entities": all_entities, "currencies": all_currencies, **multi, "expected": expected,
            "gaps": gaps, "checks": checks, "gl_periods": gl_periods, "fy_start_month": fy_start,
            "reporting_currency": next((r.get("Reporting Currency") for r in config_rows
                                        if r.get("Reporting Currency")), None)}


def write_report(path, period, prof, an, notes):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Validation Report"
    bold, hdr_fill = Font(bold=True), PatternFill("solid", fgColor="1F3864")
    fills = {"Present": "C6EFCE", "PASS": "C6EFCE", "Missing": "FFC7CE", "FAIL": "FFC7CE",
             "REVIEW": "FFEB9C", "LOW": "FFC7CE", "MEDIUM": "FFEB9C", "Not needed": "EDEDED",
             "Not provided (optional)": "EDEDED", "INFO": "DDEBF7"}
    r = 1

    def line(*vals, header=False, status_col=None):
        nonlocal r
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if header:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), hdr_fill
        if status_col and vals[status_col - 1] in fills:
            ws.cell(row=r, column=status_col).fill = PatternFill("solid", fgColor=fills[vals[status_col - 1]])
        r += 1

    def section(title):
        nonlocal r
        r += 1
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=13)
        r += 1

    ws.cell(row=r, column=1, value=f"{period} Data Validation Report").font = Font(bold=True, size=16)
    r += 1
    line(f"Inputs folder: {prof['inputs_dir']}")
    line(f"Entities: {', '.join(an['entities']) or 'none detected'}  |  Currencies: "
         f"{', '.join(an['currencies']) or 'none detected'}  |  Reporting currency (config): "
         f"{an['reporting_currency'] or 'not stated'}")
    line("Status: AWAITING USER CONFIRMATION -- nothing has been consolidated yet (Phase 2.5 checkpoint).")

    section("1. Expected vs received")
    line("Category", "Status", "Matched source (triage confidence)", header=True)
    for e in an["expected"]:
        line(e["category"], e["status"], e["matched"], status_col=2)
    for f in prof["needs_extraction"]:
        line(f, "Needs extraction", "Non-tabular file -- extracted by Claude in Phase 1, see notes below")
    for u in prof["unreadable"]:
        line(u["file"], "FAIL", f"Unreadable: {u['error']}", status_col=2)

    section("2. Coverage per source")
    line("Source", "Category", "Conf.", "Rows", "Entities", "Periods", "Currencies", "Missing contract columns",
         header=True)
    for t in prof["tables"]:
        ps = t["periods"]
        line(t["source"], t["category"] or "UNRECOGNIZED", t["confidence"], t["rows"], ", ".join(t["entities"]),
             f"{ps[0]} to {ps[-1]} ({len(ps)})" if ps else "", ", ".join(t["currencies"]),
             ", ".join(t["missing_columns"]), status_col=3)

    section("3. Cross-category gaps")
    line("Type", "Category", "Detail", header=True)
    for g in an["gaps"] or [{"type": "None found", "category": "", "detail": ""}]:
        line(g["type"], g["category"], g["detail"])

    section("4. Structural, data-shape & group-structure checks (automated; Claude adds BS/CF/headcount "
            "tie-outs below; 'playbook' = references/complex-scenarios-playbook.md)")
    line("Check", "Scope", "Result", "Detail", header=True)
    for c in an["checks"] or [{"check": "No automated issues", "scope": "", "result": "PASS", "detail": ""}]:
        line(c["check"], c["scope"], c["result"], c["detail"], status_col=3)

    section("5. Judgment items for review (LOW = needs your answer; MEDIUM = worth a glance)")
    line("Tier", "Item", "Detail", header=True)
    for n in notes or [{"tier": "", "item": "No judgment items recorded", "detail": ""}]:
        line(n.get("tier", ""), n.get("item", ""), n.get("detail", ""), status_col=1)

    for letter, width in zip("ABCDEFGH", (34, 22, 40, 10, 22, 26, 14, 40)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    return brlib.safe_save(wb.save, path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--period", help="Period label, e.g. 2026-Q1 (default: derived from GL periods)")
    ap.add_argument("--notes", help="JSON file of judgment items to include in section 5")
    ap.add_argument("--inputs", help="Override the inputs folder (default: from config/data-paths.json)")
    ap.add_argument("--outputs", help="Override the outputs folder (default: from config/data-paths.json)")
    ap.add_argument("--no-report", action="store_true", help="Profile only; don't write the report")
    args = ap.parse_args()

    inputs_dir, outputs_dir, using_drive = brlib.resolve_paths(create=True)
    inputs_dir = args.inputs or inputs_dir
    outputs_dir = args.outputs or outputs_dir
    os.makedirs(outputs_dir, exist_ok=True)
    if not using_drive:
        print("[INFO] Google Drive not configured/reachable -- using local data/ folders.")
    prof = profile(inputs_dir)
    if not prof["files"]:
        print(f"[FAIL] No input files found in {inputs_dir}")
        sys.exit(2)
    an = analyse(prof)
    period = args.period or brlib.period_label(an["gl_periods"], an["fy_start_month"])
    notes = []
    if args.notes:
        with open(args.notes, encoding="utf-8") as f:
            notes = json.load(f)

    with open(os.path.join(outputs_dir, "_input_profile.json"), "w", encoding="utf-8") as f:
        json.dump({"period": period, "profile": prof, "analysis": an}, f, indent=1, default=str)

    print(f"Inputs: {inputs_dir}  ({len(prof['files'])} file(s))")
    if len(prof["files"]) > 1:
        print("[INFO] More than one input file -- tabs pooled per category (one workbook is the standard case).")
    print(f"Period label: {period}  |  Entities: {', '.join(an['entities'])}  |  Currencies: "
          f"{', '.join(an['currencies'])}")
    for e in an["expected"]:
        print(f"  {e['status']:<11} {e['category']:<24} {e['matched']}")
    for t in prof["tables"]:
        if t["category"] is None:
            print(f"  UNRECOGNIZED {t['source']} (header coverage {t['header_coverage']:.0%})")
    for x in prof["needs_extraction"]:
        print(f"  NEEDS EXTRACTION {x}")
    for u in prof["unreadable"]:
        print(f"  UNREADABLE {u['file']}: {u['error']}")
    print(f"Gaps: {len(an['gaps'])}  |  Checks needing review/failing: "
          f"{sum(c['result'] in ('FAIL', 'REVIEW') for c in an['checks'])}")
    for g in an["gaps"]:
        print(f"  - {g['detail']}")
    for c in an["checks"]:
        if c["result"] != "PASS":
            print(f"  - [{c['result']}] {c['check']} @ {c['scope']}: {c['detail']}")
    if not args.no_report:
        written = write_report(os.path.join(outputs_dir, f"{period}_Data_Validation_Report.xlsx"),
                               period, prof, an, notes)
        print(f"Report: {written}")


if __name__ == "__main__":
    main()
