#!/usr/bin/env python3
"""
Phase 8 QA -- text-level checks across every deliverable for one period.

  python scripts/final_qa.py --period 2026-Q1 --currency USD [--period-text "Q1 FY2026"]
                             [--basis "US GAAP"] [--expect "Revenue=12,345,678" ...]

Checks, per deliverable (<Period>_*.xlsx / .docx / .pptx in the resolved outputs folder; the Data Validation
Report and the Reconciliation & Sign-off file are excluded -- the sign-off file has blank client columns by design):
  - no placeholder text (TODO, TBD, {{ }}, Lorem ipsum, ...)
  - currency code, fiscal-period text, and accounting basis are stated
  - every --expect headline figure appears verbatim in the memo AND the deck (cross-document check)
Exit code 0 = PASS, 1 = FAIL, 2 = deliverables missing.
"""
import argparse
import glob
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import brlib  # noqa: E402

PLACEHOLDERS = [r"\bTODO\b", r"\bTBD\b", r"\bFIXME\b", r"\bPLACEHOLDER\b", r"\{\{", r"\}\}", r"Lorem ipsum",
                r"<insert", r"\bXXX\b", r"\[(?:insert|add|fill)[^\]]*\]"]
BASIS_TERMS = ["GAAP", "IFRS", "Ind AS", "accounting basis", "accounting standard"]


def text_of(path):
    ext = os.path.splitext(path)[1].lower()
    parts = []
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts += [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [c.text for row in t.rows for c in row.cells]
        for s in d.sections:
            parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    elif ext == ".pptx":
        from pptx import Presentation
        for slide in Presentation(path).slides:
            for sh in slide.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    parts += [c.text for row in sh.table.rows for c in row.cells]
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            for ws in wb.worksheets:
                parts.append(ws.title)
                parts += [v for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str)]
        finally:
            wb.close()
    return "\n".join(parts)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--period", required=True, help="File prefix, e.g. 2026-Q1")
    ap.add_argument("--currency", required=True, help="Reporting currency code, e.g. USD")
    ap.add_argument("--period-text", help="How the period is written in prose (default: the --period value "
                                          "or its year)")
    ap.add_argument("--basis", action="append", help="Accounting basis text that must appear (repeatable)")
    ap.add_argument("--expect", action="append", default=[], help='Headline figure, "Label=1,234,567"')
    ap.add_argument("--outputs", help="Override the outputs folder")
    args = ap.parse_args()

    outputs = args.outputs or brlib.resolve_paths()[1]
    files = [f for ext in ("xlsx", "docx", "pptx")
             for f in glob.glob(os.path.join(outputs, f"{args.period}_*.{ext}"))
             if "Data_Validation_Report" not in f and "Reconciliation_and_Signoff" not in f
             and not os.path.basename(f).startswith("~$")]
    kinds = {os.path.splitext(f)[1] for f in files}
    missing = {".xlsx", ".docx", ".pptx"} - kinds
    if missing:
        print(f"[FAIL] Missing deliverable type(s) for {args.period}: {', '.join(sorted(missing))} in {outputs}")
        sys.exit(2)

    year = re.match(r"\d{4}", args.period)
    period_terms = [args.period_text] if args.period_text else [args.period] + ([year.group(0)] if year else [])
    basis_terms = args.basis or BASIS_TERMS
    failures, texts = [], {}

    for f in sorted(files):
        name = os.path.basename(f)
        try:
            t = texts[f] = text_of(f)
        except Exception as e:
            failures.append(f"{name}: unreadable ({type(e).__name__}: {e})")
            continue
        for pat in PLACEHOLDERS:
            for m in re.finditer(pat, t, flags=re.IGNORECASE):
                ctx = t[max(0, m.start() - 40):m.end() + 40].replace("\n", " ")
                failures.append(f"{name}: placeholder '{m.group(0)}' ... {ctx} ...")
        if not re.search(rf"\b{re.escape(args.currency)}\b", t):
            failures.append(f"{name}: currency code {args.currency} never stated")
        if not any(p.lower() in t.lower() for p in period_terms):
            failures.append(f"{name}: fiscal period ({' / '.join(period_terms)}) never stated")
        if not any(b.lower() in t.lower() for b in basis_terms):
            failures.append(f"{name}: accounting basis ({' / '.join(basis_terms)}) never stated")
        print(f"[checked] {name}: {len(t):,} chars")

    for exp in args.expect:
        label, _, value = exp.partition("=")
        for f, t in texts.items():
            if f.endswith((".docx", ".pptx")) and "Business_Review" in f and value.strip() not in t:
                failures.append(f"{os.path.basename(f)}: headline '{label}' value {value.strip()} not found "
                                f"(cross-document mismatch with workbook?)")

    print()
    for x in failures:
        print(f"[FAIL] {x}")
    print("PASS" if not failures else f"FAIL ({len(failures)} issue(s))")
    sys.exit(0 if not failures else 1)


if __name__ == "__main__":
    main()
